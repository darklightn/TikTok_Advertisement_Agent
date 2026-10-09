from __future__ import annotations

import csv
import json
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook


EMPTY_MARKERS = {"", "-", "--", "n/a", "na", "nan", "none", "null"}
ID_COLUMNS = {
    "店铺ID",
    "广告户ID",
    "计划ID",
    "商品ID",
    "Post ID",
    "SKU ID",
    "订单ID",
    "售后单号",
    "变更ID",
    "对象ID",
    "营销活动ID",
    "支付账户ID",
}
PERCENT_COLUMNS = {
    "广告CTR",
    "ROI竞争力百分位",
    "成本占计划比",
    "订单占计划比",
    "收入占计划比",
    "商品广告点击率",
    "广告转化率",
    "2秒广告视频播放率",
    "6秒广告视频播放率",
    "25%广告视频播放率",
    "50%广告视频播放率",
    "75%广告视频播放率",
    "100%广告视频播放率",
    "平均折扣率",
    "平台费率",
    "按当前售价估算毛利率",
    "商品点击率",
    "加购率",
    "CTOR",
    "退款金额率",
    "退款件数率",
    "成本覆盖率",
    "估算商品毛利率",
    "商责取消率",
    "延迟发货率",
    "商责差评率",
    "商家责任退货退款率",
    "快速发货率",
}


class CleaningError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class CleaningContext:
    shop_root: Path
    target: str
    start_date: date | None
    end_date: date | None
    output_dir: Path
    shop_id: str | None = None


@dataclass(frozen=True)
class LocatedSources:
    files: dict[str, list[Path]]
    metadata: dict[str, object]


def normalize_placeholder(value: Any) -> Any | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, str) and value.strip().lower() in EMPTY_MARKERS:
        return None
    return value


def normalize_text(value: Any) -> str | None:
    value = normalize_placeholder(value)
    if value is None:
        return None
    text = str(value).replace("\t", "").strip()
    return text or None


def select_official_overview_total_row(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Return the unique explicit '-' total; never sum detail rows as a fallback."""
    dimensions = {field for field in ("时间", "按天") if rows and field in rows[0]}
    if not dimensions:
        raise CleaningError("INVALID_OVERVIEW_SCHEMA", "GMV Max计划总览缺少时间或按天字段")
    totals: list[dict[str, Any]] = []
    for row in rows:
        if {field for field in ("时间", "按天") if field in row} != dimensions:
            raise CleaningError("INVALID_OVERVIEW_SCHEMA", "GMV Max计划总览时间维度字段不一致")
        # Check the explicit marker before normalization erases '-' along with
        # unrelated missing values such as None, blank strings and N/A.
        markers = [isinstance(row[field], str) and row[field].strip() == "-" for field in dimensions]
        if any(markers) and not all(markers):
            raise CleaningError("INVALID_OVERVIEW_SCHEMA", "GMV Max计划总览时间维度总计标记冲突")
        if all(markers):
            totals.append(row)
    if not totals:
        raise CleaningError("MISSING_OVERVIEW_TOTAL", "GMV Max计划总览缺少官方总计行")
    if len(totals) != 1:
        raise CleaningError("DUPLICATE_OVERVIEW_TOTAL", "GMV Max计划总览存在多条官方总计行")
    return totals[0]


def normalize_id(value: Any) -> str | None:
    value = normalize_placeholder(value)
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value)


def normalize_date(value: Any) -> date | None:
    value = normalize_placeholder(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = pd.to_datetime(str(value).strip(), errors="coerce")
    if pd.isna(parsed):
        raise CleaningError("INVALID_DATE", f"无法解析日期：{value}")
    return parsed.date()


def normalize_datetime(value: Any) -> datetime | None:
    value = normalize_placeholder(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        parsed_datetime = value
        if parsed_datetime.tzinfo is not None:
            return parsed_datetime.astimezone(timezone.utc).replace(tzinfo=None)
        return parsed_datetime
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    parsed = pd.to_datetime(str(value).strip(), errors="coerce")
    if pd.isna(parsed):
        raise CleaningError("INVALID_DATETIME", f"无法解析日期时间：{value}")
    parsed_datetime = parsed.to_pydatetime()
    if parsed_datetime.tzinfo is not None:
        return parsed_datetime.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed_datetime


def normalize_percentage_text(value: Any) -> float | None:
    value = normalize_placeholder(value)
    if value is None:
        return None
    text = normalize_text(value)
    if text is None:
        return None
    if not text.endswith("%"):
        raise CleaningError("INVALID_PERCENTAGE", f"百分比缺少%：{value}")
    number = normalize_money(text[:-1])
    return None if number is None else float(number) / 100


def normalize_hour(value: Any) -> str | None:
    value = normalize_placeholder(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%H:00")
    if isinstance(value, time):
        return value.strftime("%H:00")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        hour = int(value)
        if float(value) == hour and 0 <= hour <= 23:
            return f"{hour:02d}:00"
    text = str(value).strip()
    match = re.fullmatch(r"(\d{1,2})(?::(\d{1,2}))?(?::\d{1,2})?", text)
    if not match:
        raise CleaningError("INVALID_HOUR", f"无法解析小时：{value}")
    hour = int(match.group(1))
    if not 0 <= hour <= 23:
        raise CleaningError("INVALID_HOUR", f"小时超出范围：{value}")
    return f"{hour:02d}:00"


def normalize_money(value: Any) -> float | int | None:
    value = normalize_placeholder(value)
    if value is None:
        return None
    if isinstance(value, bool):
        raise CleaningError("INVALID_NUMBER", f"无法解析金额：{value}")
    if isinstance(value, (int, float, Decimal)):
        number = float(value)
    else:
        cleaned = re.sub(r"[^0-9.\-]", "", str(value).replace(",", ""))
        if cleaned in {"", "-", ".", "-."}:
            return None
        try:
            number = float(Decimal(cleaned))
        except (InvalidOperation, ValueError) as exc:
            raise CleaningError("INVALID_NUMBER", f"无法解析金额：{value}") from exc
    return int(number) if number.is_integer() else number


def normalize_integer(value: Any) -> int | None:
    number = normalize_money(value)
    if number is None:
        return None
    if float(number).is_integer():
        return int(number)
    raise CleaningError("INVALID_INTEGER", f"无法解析整数：{value}")


def safe_divide(numerator: Any, denominator: Any, multiplier: float = 1.0) -> float | None:
    numerator = normalize_placeholder(numerator)
    denominator = normalize_placeholder(denominator)
    if numerator is None or denominator is None:
        return None
    denominator_number = float(denominator)
    if denominator_number == 0:
        return None
    return float(numerator) / denominator_number * multiplier


def parse_json_object(value: Any, field_name: str = "JSON") -> dict[str, Any]:
    value = normalize_placeholder(value)
    if value is None:
        raise CleaningError("INVALID_JSON", f"{field_name}为空")
    if isinstance(value, dict):
        return value
    try:
        parsed = json.loads(str(value))
    except (TypeError, json.JSONDecodeError) as exc:
        raise CleaningError("INVALID_JSON", f"{field_name}解析失败") from exc
    if not isinstance(parsed, dict):
        raise CleaningError("INVALID_JSON", f"{field_name}不是JSON对象")
    return parsed


def parse_json_value(value: Any, field_name: str = "JSON") -> Any:
    value = normalize_placeholder(value)
    if value is None:
        raise CleaningError("INVALID_JSON", f"{field_name}为空")
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(str(value))
    except (TypeError, json.JSONDecodeError) as exc:
        raise CleaningError("INVALID_JSON", f"{field_name}解析失败") from exc


def _rows_from_values(
    values: Iterable[tuple[Any, ...]],
    path: Path,
    required_headers: set[str] | None = None,
) -> list[dict[str, Any]]:
    materialized = list(values)
    if not materialized:
        raise CleaningError("EMPTY_SOURCE", f"原始文件为空：{path.name}")
    header_index = None
    for index, row in enumerate(materialized[:25]):
        populated = [normalize_text(value) for value in row]
        populated_set = {value for value in populated if value is not None}
        if required_headers is not None:
            is_header = required_headers.issubset(populated_set)
        else:
            is_header = len(populated_set) >= 2
        if is_header:
            header_index = index
            break
    if header_index is None:
        raise CleaningError("MISSING_HEADER", f"找不到表头：{path.name}")
    headers = [normalize_text(value) for value in materialized[header_index]]
    if not headers or any(header is None for header in headers):
        raise CleaningError("INVALID_HEADER", f"表头包含空列：{path.name}")
    if len(set(headers)) != len(headers):
        raise CleaningError("DUPLICATE_HEADER", f"表头包含重复字段：{path.name}")
    rows: list[dict[str, Any]] = []
    for raw in materialized[header_index + 1 :]:
        if all(normalize_placeholder(value) is None for value in raw):
            continue
        padded = list(raw) + [None] * (len(headers) - len(raw))
        rows.append(dict(zip(headers, padded[: len(headers)])))
    return rows


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise CleaningError("INVALID_XLSX", f"无法读取Excel：{path.name}") from exc
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        # TikTok 导出的部分 XLSX 存在错误的 worksheet dimension。
        # read_only 模式会信任该元数据，因此必须重置后按真实单元格重新扫描。
        worksheet.reset_dimensions()
        return _rows_from_values(worksheet.iter_rows(values_only=True), path)
    finally:
        workbook.close()


def read_xlsx_rows_with_header(path: Path, required_headers: set[str]) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise CleaningError("INVALID_XLSX", f"无法读取Excel：{path.name}") from exc
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        worksheet.reset_dimensions()
        return _rows_from_values(worksheet.iter_rows(values_only=True), path, required_headers)
    finally:
        workbook.close()


def read_product_analytics_rows(path: Path) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise CleaningError("INVALID_XLSX", f"无法读取Excel：{path.name}") from exc
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        worksheet.reset_dimensions()
        values = list(worksheet.iter_rows(values_only=True))
        header_index = None
        for index, row in enumerate(values[:25]):
            fields = {normalize_text(value) for value in row}
            if {"商品名", "商品 ID", "GMV"}.issubset(fields):
                header_index = index
                break
        if header_index is None or header_index == 0:
            raise CleaningError("MISSING_HEADER", f"找不到商品分析多级表头：{path.name}")

        upper = list(values[header_index - 1])
        lower = list(values[header_index])
        width = len(lower)
        upper.extend([None] * (width - len(upper)))
        headers: list[str] = []
        current_group: str | None = None
        plain_fields = {"商品名", "商品 ID", "GMV 区间", "发品状态"}
        for column, raw_field in enumerate(lower):
            field = normalize_text(raw_field)
            if field is None:
                raise CleaningError("INVALID_HEADER", f"商品分析表头包含空列：{path.name}")
            group = normalize_text(upper[column])
            if group is not None:
                current_group = group
            headers.append(field if field in plain_fields or current_group is None else f"{current_group}.{field}")
        if len(set(headers)) != len(headers):
            raise CleaningError("DUPLICATE_HEADER", f"商品分析多级表头仍有重复字段：{path.name}")

        rows: list[dict[str, Any]] = []
        for raw in values[header_index + 1 :]:
            if all(normalize_placeholder(value) is None for value in raw):
                continue
            padded = list(raw) + [None] * (width - len(raw))
            rows.append(dict(zip(headers, padded[:width])))
        return rows
    finally:
        workbook.close()


def read_xlsx_custom_properties(path: Path) -> dict[str, str]:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise CleaningError("INVALID_XLSX", f"无法读取Excel：{path.name}") from exc
    try:
        return {
            str(item.name): str(item.value)
            for item in workbook.custom_doc_props
            if getattr(item, "name", None) is not None and getattr(item, "value", None) is not None
        }
    finally:
        workbook.close()


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
    except UnicodeDecodeError as exc:
        raise CleaningError("INVALID_CSV", f"CSV不是UTF-8编码：{path.name}") from exc
    return _rows_from_values((tuple(row) for row in rows), path)


def read_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return read_xlsx_rows(path)
    if suffix == ".csv":
        return read_csv_rows(path)
    raise CleaningError("UNSUPPORTED_SOURCE", f"不支持的原始文件类型：{path.name}")


def read_account_identity(path: Path) -> dict[str, Any]:
    rows = read_rows(path)
    if not rows:
        raise CleaningError("EMPTY_SOURCE", "广告账户状态文件没有数据行")
    advertiser_ids = {normalize_id(row.get("advertiser_id")) for row in rows}
    advertiser_ids.discard(None)
    if len(advertiser_ids) != 1:
        raise CleaningError("IDENTITY_CONFLICT", "广告账户状态无法确定唯一广告户")
    row = rows[0]
    snapshot = parse_json_object(row.get("account_snapshot_json"), "account_snapshot_json")
    frozen = snapshot.get("frozen_identity") if isinstance(snapshot.get("frozen_identity"), dict) else {}
    payment = snapshot.get("payment") if isinstance(snapshot.get("payment"), dict) else {}
    adv_info = payment.get("adv_info_vo") if isinstance(payment.get("adv_info_vo"), dict) else {}
    return {
        "shop_id": normalize_id(frozen.get("seller_id")),
        "advertiser_id": next(iter(advertiser_ids)),
        "advertiser_name": normalize_text(row.get("advertiser_name")),
        "timezone": normalize_text(adv_info.get("timezone")),
        "currency": normalize_text(row.get("currency")) or normalize_text(adv_info.get("currency")),
    }


def write_formatted_xlsx(frame: pd.DataFrame, path: Path, columns: list[str]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "数据"
    worksheet.append(columns)
    for record in frame.reindex(columns=columns).itertuples(index=False, name=None):
        worksheet.append([None if pd.isna(value) else value for value in record])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, name in enumerate(columns, start=1):
        if name in ID_COLUMNS:
            for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    if item.value is not None:
                        item.value = str(item.value)
                    item.number_format = "@"
        elif name in PERCENT_COLUMNS:
            for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    item.number_format = "0.00%"
        elif name in {"当前值", "较上期变化", "上期值", "行业均值"}:
            for row_index, record in enumerate(frame.reindex(columns=columns).to_dict(orient="records"), start=2):
                unit = normalize_text(record.get("单位"))
                metric = normalize_text(record.get("指标"))
                target = normalize_text(record.get("目标条件"))
                is_percent = unit == "%" or (
                    name == "当前值" and (bool(metric and metric.endswith("率")) or bool(target and "%" in target))
                )
                if is_percent:
                    worksheet.cell(row=row_index, column=index).number_format = "0.00%"
        elif name.endswith("日期") or name in {"日期", "统计开始日期", "统计结束日期"}:
            for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    item.number_format = "yyyy-mm-dd"
        elif name.endswith("时间"):
            for cell in worksheet.iter_cols(min_col=index, max_col=index, min_row=2):
                for item in cell:
                    item.number_format = "yyyy-mm-dd hh:mm:ss"
    workbook.save(path)
