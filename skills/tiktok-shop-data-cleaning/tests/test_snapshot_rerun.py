from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
RUNNER = SKILL_ROOT / "scripts" / "run_cleaning.py"
SHOP_ID = "7496176775586483191"


def _write_xlsx(path: Path, headers: list[str], rows: list[list[Any]], sheet: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)


def _run(shop_root: Path, target: str) -> subprocess.CompletedProcess[str]:
    environment = os.environ.copy()
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    environment["PYTHONUTF8"] = "1"
    return subprocess.run(
        [
            sys.executable,
            str(RUNNER),
            "--shop-root",
            str(shop_root),
            "--shop-id",
            SHOP_ID,
            "--target",
            target,
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=environment,
    )


def _rows(path: Path) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, data_only=True)
    try:
        return list(workbook.active.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()


def _assert_identical_rerun(shop_root: Path, target: str) -> None:
    first = _run(shop_root, target)
    assert first.returncode == 0, first.stdout + first.stderr
    output = shop_root / "20_清洗数据" / f"{target}.xlsx"
    first_rows = _rows(output)

    second = _run(shop_root, target)
    assert second.returncode == 0, second.stdout + second.stderr
    assert json.loads(second.stdout)["status"] == "success"
    assert _rows(output) == first_rows


def test_shop_health_identical_snapshot_rerun(tmp_path: Path) -> None:
    shop_root = tmp_path / "测试店铺"
    headers = ["采集时间", "类别", "指标", "当前值", "目标条件", "评估开始", "评估结束", "违规次数", "说明"]
    _write_xlsx(
        shop_root / "店铺数据" / "店铺健康快照" / "2026-09-04.xlsx",
        headers,
        [["2026-09-04T05:15:22.082Z", "履约", "快速发货率", "0.19402824188854153%", ">= 90%", "2026-07-25", "2026-08-23", 0, None]],
        "店铺健康诊断快照",
    )
    _assert_identical_rerun(shop_root, "店铺健康明细")


def test_fulfillment_identical_snapshot_rerun(tmp_path: Path) -> None:
    shop_root = tmp_path / "测试店铺"
    headers = [
        "采集时间", "类别", "指标", "当前值", "单位", "较上期变化", "上期值", "行业均值", "行业分类",
        "评估开始", "评估结束", "指标日期", "分子", "分母", "官方状态", "统计口径",
    ]
    _write_xlsx(
        shop_root / "店铺数据" / "履约表现快照" / "2026-09-04.xlsx",
        headers,
        [["2026-09-04T05:14:09.528Z", "履约表现指标", "快速发货率", 98.37, "%", 1.27, 97.1, 95.5, "Lifestyle", "2026-07-25", "2026-08-23", "2026-08-30", 9996, 10162, "良好", "官方7天窗口"]],
        "履约表现诊断快照",
    )
    _assert_identical_rerun(shop_root, "履约表现明细")


def test_refund_empty_logistics_time_identical_snapshot_rerun(tmp_path: Path) -> None:
    shop_root = tmp_path / "测试店铺"
    order_path = shop_root / "店铺数据" / "订单明细" / "2026-08-01_至_2026-08-31.csv"
    order_path.parent.mkdir(parents=True, exist_ok=True)
    with order_path.open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerow(["Order ID", "Order Amount"])

    headers = [
        "采集时间", "订单ID", "售后单号", "售后申请时间", "售后类型", "售后原因", "客户说明", "商品摘要",
        "订单实付", "退款金额", "币种", "处理状态", "退货物流状态", "退货物流时间", "平台标签", "客户凭证数量",
    ]
    _write_xlsx(
        shop_root / "店铺数据" / "退货退款快照" / "2026-09-04.xlsx",
        headers,
        [
            ["2026-09-04T05:15:23.837Z", "ORDER-1", "AFTER-1", "2026-09-01T11:55:46.875Z", "退货退款", "商品与描述不符", None, "1 件商品", "已付款", "₱155.99", "PHP", "退款完成", "已送达", "2026-09-03T04:14:24.000Z", None, 1],
            ["2026-09-04T05:15:23.837Z", "ORDER-2", "AFTER-2", "2026-09-01T12:00:00.000Z", "退货退款", "不再需要", None, "1 件商品", "已付款", "₱99.00", "PHP", "处理中", None, None, None, 0],
        ],
        "退货退款诊断快照",
    )
    _assert_identical_rerun(shop_root, "退货退款明细")
