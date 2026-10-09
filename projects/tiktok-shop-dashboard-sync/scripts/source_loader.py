"""Read selected shops and verified in-memory XLSX snapshots; never write upstream."""
from datetime import date, datetime
from io import BytesIO
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook


def issue(code, shop, source=None, period=None, message=None):
    return {'code': code, 'message': message or code, 'shop_directory': shop,
            'source': source, 'period': period}


def empty(value):
    return value is None or isinstance(value, str) and not value.strip()


def text_id(value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError('身份必须为非空文本，不能从数字恢复ID')
    return value.strip()


def business_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and re.fullmatch(r'[0-9]{4}-[0-9]{1,2}-[0-9]{1,2}', value.strip()):
        return date(*map(int, value.strip().split('-'))).isoformat()
    raise ValueError('无效业务日期')


def discover(data_root, scope):
    if not scope['all_shops']:
        return scope['shops']
    return sorted(p.name for p in data_root.iterdir() if p.is_dir() and (p / '20_清洗数据').is_dir())


def read_mappings(directory, definitions):
    matches = {source: [] for source in definitions}
    if not directory.is_dir():
        return matches
    for path in sorted(directory.iterdir()):
        if not re.fullmatch(r'V[0-9]{3,}\.json', path.name) or not path.is_file():
            continue
        obj = json.loads(path.read_text(encoding='utf-8-sig'))
        if not isinstance(obj, dict) or not isinstance(obj.get('tables'), list):
            raise ValueError('分卷根结构或tables类型无效')
        for item in obj['tables']:
            # Only relevant source mappings need a valid table ID and hash.
            if not isinstance(item, dict):
                continue
            for source, definition in definitions.items():
                if item.get('local_path') == definition['relative_path']:
                    matches[source].append((path.resolve(), item))
    return matches


def parse_snapshot(content, definition, shop, source):
    issues, rows = [], []
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if definition['sheet'] not in workbook.sheetnames:
            raise ValueError('缺少工作表：' + definition['sheet'])
        iterator = workbook[definition['sheet']].iter_rows(values_only=True)
        headers = list(next(iterator, ()))
        required = definition['required_headers']
        if any(headers.count(name) != 1 for name in required):
            raise ValueError('必需表头缺失或重复')
        indexes = {name: headers.index(name) for name in required}
        for line, values in enumerate(iterator, 2):
            if all(empty(v) for v in values):
                continue
            row = {name: values[index] if index < len(values) else None for name, index in indexes.items()}
            period = None
            try:
                start, end = (business_date(row[k]) for k in ['统计开始日期', '统计结束日期'])
                if start > end:
                    raise ValueError('来源起止日期倒置')
                period = {'start_date': start, 'end_date': end}
                row['统计开始日期'], row['统计结束日期'] = start, end
                for name in ['店铺ID'] + (['商品ID'] if source == 'product' else ['广告户ID'] if source == 'ads' else []):
                    row[name] = text_id(row[name])
            except ValueError as exc:
                issues.append(issue('SOURCE_IDENTITY_INVALID', shop, source, period, f'第{line}行：{exc}'))
                continue
            try:
                row['店铺名称'] = text_id(row['店铺名称'])
            except ValueError:
                issues.append(issue('SOURCE_NAME_CONFLICT', shop, source, None, f'第{line}行店铺名称不明确'))
                continue
            row['_line'] = line
            rows.append(row)
    except Exception as exc:
        # Workbook/ZIP/XML decoding errors are source errors, not a zero-row dataset.
        issues.append(issue('SOURCE_SCHEMA_INVALID', shop, source, message='工作簿结构无法读取：' + str(exc)))
        rows = []
    finally:
        if workbook is not None:
            workbook.close()
    return rows, issues


def load_shop(data_root, upstream_root, shop, definitions):
    directory = (data_root / shop).resolve()
    result = {'shop': shop, 'rows': {}, 'issues': [], 'source_files': []}
    if directory.parent != data_root or not directory.is_dir():
        result['issues'].append(issue('SHOP_NOT_FOUND', shop, message='指定店铺不是来源根目录中的有效直接子目录'))
        return result
    upstream = (upstream_root / shop).resolve()
    if upstream.parent != upstream_root:
        result['issues'].append(issue('UPSTREAM_STATE_INVALID', shop, message='上游店铺路径超出直接子目录'))
        return result
    try:
        matches = read_mappings(upstream, definitions)
    except (OSError, ValueError) as exc:
        result['issues'].append(issue('UPSTREAM_STATE_INVALID', shop, message='上游分卷无效：' + str(exc)))
        return result
    snapshots = {}
    for source, definition in definitions.items():
        path = (directory / definition['relative_path']).resolve()
        if not path.is_relative_to(directory) or not path.is_file():
            result['issues'].append(issue('SOURCE_FILE_MISSING', shop, source, message='缺少必需来源文件'))
            continue
        entries = matches[source]
        if len(entries) != 1:
            code = 'UPSTREAM_MAPPING_MISSING' if not entries else 'UPSTREAM_MAPPING_CONFLICT'
            result['issues'].append(issue(code, shop, source, message='来源需要唯一的精确路径映射'))
            continue
        state_file, mapping = entries[0]
        digest = mapping.get('file_hash')
        if (not isinstance(mapping.get('table_id'), str) or not mapping['table_id'].strip()
                or not isinstance(digest, str) or not re.fullmatch(r'[a-fA-F0-9]{64}', digest)):
            result['issues'].append(issue('UPSTREAM_STATE_INVALID', shop, source, message='映射table_id或file_hash无效'))
            continue
        try:
            content = path.read_bytes()
        except OSError as exc:
            result['issues'].append(issue('SOURCE_FILE_UNREADABLE', shop, source, message=str(exc)))
            continue
        actual = hashlib.sha256(content).hexdigest()
        if actual != digest.lower():
            result['issues'].append(issue('SOURCE_VERSION_MISMATCH', shop, source, message='本地快照Hash与上游登记不一致'))
            continue
        result['source_files'].append({'shop_directory': shop, 'source': source, 'path': str(path),
            'file_hash': actual, 'upstream_file_hash': digest, 'state_file': str(state_file), 'table_id': mapping['table_id']})
        snapshots[source] = content
    # Hash verification is completed for all possible sources before schema/identity checks.
    for source, content in snapshots.items():
        rows, problems = parse_snapshot(content, definitions[source], shop, source)
        result['rows'][source] = rows
        result['issues'].extend(problems)
    return result
