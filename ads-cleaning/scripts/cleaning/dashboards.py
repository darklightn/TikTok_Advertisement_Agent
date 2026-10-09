"""Build human-facing, single-sheet dashboard workbooks from one cleaning batch."""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import sys
from typing import Any, Iterable
import zipfile

_RUNTIME = Path(__file__).resolve().parents[2] / ".runtime"
if _RUNTIME.is_dir():
    sys.path.insert(0, str(_RUNTIME))


class DashboardBuildError(RuntimeError):
    pass


DASHBOARD_CONTRACT_VERSION = "2026-08-14.5"


_XLSX_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


def _save_deterministic_xlsx(workbook: Any, target: Path) -> None:
    workbook.properties.created = datetime(1980, 1, 1)
    workbook.properties.modified = datetime(1980, 1, 1)
    temporary = target.with_suffix(".nondeterministic.xlsx")
    workbook.save(temporary)
    with zipfile.ZipFile(temporary, "r") as source, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_STORED
    ) as destination:
        for name in source.namelist():
            original = source.getinfo(name)
            info = zipfile.ZipInfo(name, _XLSX_TIMESTAMP)
            info.compress_type = zipfile.ZIP_STORED
            info.external_attr = original.external_attr
            info.internal_attr = original.internal_attr
            info.flag_bits = original.flag_bits
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                    rb"\g<1>1980-01-01T00:00:00Z\g<2>",
                    content,
                )
            destination.writestr(info, content, compress_type=zipfile.ZIP_STORED)
    temporary.unlink()


ALIASES: dict[str, tuple[str, ...]] = {
    "as_of_utc": ("captured_at_utc", "captured_at"), "account_currency": ("currency",),
    "ad_currency": ("currency",), "commerce_currency": ("currency",),
    "advertiser_timezone": ("timezone",), "is_primary_advertiser": ("primary_managed_raw",),
    "is_owned_by_user": ("is_owned_by_user_raw",), "is_owned_by_business_center": ("is_owned_by_business_center_raw",),
    "billing_method_raw": ("billing_business_center_pay_method_raw",),
    "authorization_binding_status": ("authorization_operation_source_raw",),
    "authorization_capability_status": ("account_delivery_status_capability_status",),
    "campaign_name_as_of": ("campaign_name_observed", "campaign_name"),
    "campaign_name_observed_at_utc": ("captured_at_utc",), "ad_cost": ("cost",),
    "ad_gross_revenue": ("gross_revenue",), "cpo": ("cost_per_order_official",),
    "roi": ("roi_official",), "sns_opt_roi": ("sns_opt_roi_official",),
    "revenue_per_sku_order": ("revenue_per_sku_order_calculated",),
    "ctr_ratio": ("ctr_calculated", "product_ad_click_rate_official", "overall_click_rate_official"),
    "cvr_ratio": ("ad_conversion_rate_official",), "cpc": ("cpc_calculated",), "cpm": ("cpm_calculated",),
    "hour_start_local": ("metric_hour",), "configuration_as_of_utc": ("captured_at_utc",),
    "status_as_of_utc": ("captured_at_utc",), "budget_mode_raw": ("budget_mode_raw", "campaign_budget_mode_raw"),
    "primary_status_raw": ("campaign_primary_status_raw",), "opt_status_raw": ("campaign_opt_status_raw",),
    "schedule_start_time": ("start_time_raw", "schedule_start_time_raw"),
    "schedule_end_time": ("end_time_raw", "schedule_end_time_raw"),
    "eligible_status_raw": ("campaign_eligible_status_raw",), "eligible_reject_reason_raw": ("campaign_eligible_reject_reason_raw",),
    "eligible_reference_roi": ("campaign_eligible_roi",), "budget_utilization_status_raw": ("campaign_budget_utilization_status_raw",),
    "today_utilization_threshold": ("campaign_budget_utilization_today_threshold",),
    "yesterday_utilization_threshold": ("campaign_budget_utilization_yesterday_threshold",),
    "promotion_roas_multiplier": ("promotion_days_roas_bid_multiplier",),
    "promotion_budget_multiplier": ("promotion_days_budget_multiplier",),
    "auto_increase_effective_budget": ("auto_increase_budget_effective_budget",),
    "actor": ("user",), "product_name_as_of": ("product_name_observed", "product_name"),
    "creative_key": ("post_id",), "impressions": ("product_ad_impressions", "overall_impressions"),
    "clicks": ("product_ad_clicks", "overall_clicks"), "view_rate_2s_ratio": ("two_second_view_rate_official",),
    "view_rate_6s_ratio": ("six_second_view_rate_official",), "view_rate_25_ratio": ("twenty_five_percent_view_rate_official",),
    "view_rate_50_ratio": ("fifty_percent_view_rate_official",), "view_rate_75_ratio": ("seventy_five_percent_view_rate_official",),
    "view_rate_100_ratio": ("one_hundred_percent_view_rate_official",),
    "product_gmv": ("overall_gmv",), "product_attributed_orders": ("overall_orders",), "sku_orders": ("overall_sku_orders", "sku_orders"),
    "units_sold": ("overall_units_sold",), "estimated_customers_product": ("overall_estimated_customers",),
    "add_to_cart": ("overall_add_to_cart",), "unique_impressions": ("overall_unique_impressions",),
    "unique_clicks": ("overall_unique_clicks",), "unique_atc_users": ("overall_unique_atc_users",),
    "refund_amount": ("overall_refund_amount",), "refunded_units": ("overall_refunded_units",),
    "refund_customers_product": ("overall_refund_customers",), "merchant_live_gmv": ("overall_merchant_live_gmv",),
    "merchant_video_gmv": ("overall_merchant_video_gmv",), "creator_affiliate_gmv": ("overall_creator_attributed_gmv",),
    "merchant_product_card_gmv": ("overall_merchant_product_card_gmv",),
    "product_revenue_per_sku_order": ("overall_average_order_value",), "atc_rate_ratio": ("overall_add_to_cart_rate_official",),
    "ctor_ratio": ("overall_ctor_official",), "unique_ctr_ratio": ("overall_unique_click_rate_official",),
    "unique_atc_rate_ratio": ("overall_unique_atc_rate_official",), "unique_conversion_ratio": ("overall_unique_click_conversion_rate_official",),
    "inventory_as_of_utc": ("captured_at",), "warehouse_names": ("warehouse_names_observed",),
    "effective_price": ("effective_price_calculated",), "listing_status": ("listing_status_observed", "listing_status"),
    "is_visible": ("is_visible_raw",), "commercial_status_as_of_utc": ("captured_at_utc",),
    "variation": ("variation_json", "variation"), "promotion_ids_json": ("sku_level_promotion_ids_json",),
    "promotion_names_json": ("sku_level_promotion_names_observed_json",), "promotion_windows_json": ("sku_level_promotion_windows_json",),
    "metric_name_official": ("metric_name",), "as_of_utc": ("captured_at", "captured_at_utc"),
    "assessment_start_date": ("assessment_start",), "assessment_end_date": ("assessment_end",),
    "current_value_numeric": ("metric_value_official",), "current_value_unit": ("unit",),
    "previous_value": ("previous_value_official",), "period_change": ("period_change_official",),
    "industry_average": ("industry_average_official",), "risk_domain": ("category",),
}

HEADER_OVERRIDES = {
    "advertiser_id":"广告户ID", "advertiser_name":"广告户名称", "campaign_id":"计划ID", "campaign_name":"计划名称",
    "campaign_name_as_of":"计划名称（截至快照）", "campaign_name_observed_at_utc":"计划名称观测时间（UTC）",
    "product_id":"商品ID", "product_name":"商品名称", "product_name_as_of":"商品名称（截至快照）", "sku_id":"SKU ID",
    "seller_sku":"商家SKU", "metric_date":"数据日期", "metric_hour":"小时", "hour_start_local":"本地小时开始时间",
    "as_of_utc":"快照时间（UTC）", "configuration_as_of_utc":"配置快照时间（UTC）", "status_as_of_utc":"状态快照时间（UTC）",
    "inventory_as_of_utc":"库存快照时间（UTC）", "commercial_status_as_of_utc":"商品状态快照时间（UTC）",
    "ad_currency":"广告币种", "commerce_currency":"店铺币种", "account_currency":"账户币种", "billing_currency":"结算币种",
    "ad_cost":"广告消耗", "ad_gross_revenue":"广告成交金额", "sns_opt_gross_revenue":"优化成交金额",
    "impressions":"曝光次数", "clicks":"点击次数", "sku_orders":"SKU订单数", "units_sold":"售出件数",
    "ctr_ratio":"点击率", "cvr_ratio":"转化率", "cpc":"单次点击成本", "cpm":"千次曝光成本", "cpo":"单次下单成本",
    "roi":"投入产出比", "sns_opt_roi":"优化投入产出比", "revenue_per_sku_order":"每SKU订单成交金额",
    "product_gmv":"商品成交金额", "shop_gmv":"店铺成交金额", "gmv_share":"成交金额占比", "gmv_rank":"成交金额排名",
    "refund_amount":"退款金额", "returned_sku_quantity":"退回SKU数量", "refunded_units":"退款件数",
    "total_stock":"总库存", "available_stock":"可售库存", "locked_stock":"锁定库存", "sales_30d":"近30天销量",
    "forecast_30d":"未来30天预测销量", "recommended_replenishment":"建议补货量", "days_of_supply_official":"官方可售天数",
    "days_of_supply_calculated":"计算可售天数", "inventory_risk_status":"库存风险状态", "risk_reasons_json":"风险原因",
    "data_status":"数据状态", "quality_flags_json":"质量标记", "lineage_manifest_id":"来源Manifest ID",
    "is_owned_by_user":"是否由当前用户拥有", "is_owned_by_business_center":"是否由商务中心拥有", "budget":"预算",
    "target_roi":"目标投入产出比", "roas_bid":"ROAS出价", "eligible_reference_roi":"准入参考投入产出比",
    "object_id":"对象ID", "post_id":"帖子ID", "tiktok_account":"TikTok账号", "gmv_tier":"GMV层级",
    "add_to_cart":"加购次数", "unique_impressions":"去重曝光人数", "unique_clicks":"去重点击人数",
    "unique_atc_users":"去重加购人数", "merchant_live_gmv":"商家直播成交金额", "merchant_video_gmv":"商家视频成交金额",
    "creator_affiliate_gmv":"达人联盟成交金额", "sku_id":"SKU编号", "is_visible":"是否可见",
    "statistic_basis":"统计口径", "detected_at_utc":"检测时间（UTC）", "entity_id":"实体ID",
    "candidate_reasons_json":"候选原因", "reconciliation_delta":"勾稽差额",
}

TOKENS = {
    "id":"ID", "name":"名称", "raw":"原始值", "status":"状态", "date":"日期", "time":"时间", "utc":"UTC",
    "current":"当前", "value":"数值", "unit":"单位", "count":"数量", "amount":"金额", "ratio":"比例", "rate":"比率",
    "total":"总计", "average":"平均", "min":"最小", "max":"最大", "official":"官方", "calculated":"计算值",
    "product":"商品", "order":"订单", "campaign":"计划", "creative":"素材", "shop":"店铺", "inventory":"库存",
    "risk":"风险", "anomaly":"异常", "evidence":"证据", "source":"来源", "currency":"币种", "cost":"消耗",
    "revenue":"成交金额", "gross":"毛", "channel":"渠道", "customer":"客户", "visitor":"访客", "seller":"商家",
    "platform":"平台", "shipping":"运费", "discount":"优惠", "tax":"税费", "refund":"退款", "paid":"已付款",
    "completed":"已完成", "shipped":"已发货", "cancelled":"已取消", "other":"其他", "unknown":"未知",
    "with":"有", "without":"无", "created":"创建", "distinct":"去重", "effective":"生效", "snapshot":"快照",
    "listing":"上架", "available":"可售", "locked":"锁定", "promotion":"促销", "active":"活跃", "low":"低",
    "stock":"库存", "severity":"严重程度", "domain":"领域", "code":"代码", "type":"类型", "summary":"摘要",
    "operator":"运算符", "threshold":"阈值", "baseline":"基线", "window":"窗口", "start":"开始", "end":"结束",
    "resolution":"处理", "rule":"规则", "version":"版本", "decision":"决策", "required":"需要", "mapping":"映射",
    "parse":"解析", "join":"关联", "primary":"主要", "relation":"关系", "country":"国家", "timezone":"时区",
    "precision":"精度", "role":"角色", "payment":"支付", "billing":"结算", "authorization":"授权",
    "event":"事件", "actor":"操作人", "old":"旧", "new":"新", "field":"字段", "details":"详情",
    "observed":"观测", "denominator":"分母", "numerator":"分子", "previous":"上期", "period":"周期",
    "industry":"行业", "explanation":"说明", "breach":"是否越线", "finite":"有限", "price":"价格",
    "list":"标价", "sale":"售价", "warehouse":"仓库", "variation":"规格", "recommendation":"建议",
}


def visible_header(field: str) -> str:
    if field in HEADER_OVERRIDES:
        return HEADER_OVERRIDES[field]
    words = [TOKENS.get(part, part.upper() if len(part) <= 4 else f"[{part}]") for part in field.split("_")]
    return "".join(words)


def safe_div(numerator: Any, denominator: Any) -> float | None:
    try:
        if numerator in (None, "") or denominator in (None, "") or float(denominator) == 0:
            return None
        return float(numerator) / float(denominator)
    except (TypeError, ValueError):
        return None


def _value(row: dict[str, Any], field: str) -> Any:
    candidates = (field,) + ALIASES.get(field, ())
    for candidate in candidates:
        if candidate in row and row[candidate] not in (None, ""):
            value = row[candidate]
            if field == "hour_start_local" and re.fullmatch(r"\d{1,2}", str(value)):
                return f"{int(value):02d}:00"
            return value
    formulas = {
        "cpc": ("ad_cost", "clicks"), "cpm": ("ad_cost", "impressions"), "cpo": ("ad_cost", "sku_orders"),
        "roi": ("ad_gross_revenue", "ad_cost"), "revenue_per_sku_order": ("ad_gross_revenue", "sku_orders"),
        "product_revenue_per_sku_order": ("product_gmv", "sku_orders"), "ctr_ratio": ("clicks", "impressions"),
        "atc_rate_ratio": ("add_to_cart", "clicks"), "unique_ctr_ratio": ("unique_clicks", "unique_impressions"),
        "unique_atc_rate_ratio": ("unique_atc_users", "unique_clicks"),
    }
    if field in formulas:
        numerator, denominator = formulas[field]
        result = safe_div(_value(row, numerator), _value(row, denominator))
        return result * 1000 if field == "cpm" and result is not None else result
    if field == "name_resolution_status":
        return "resolved" if _value(row, "campaign_name_as_of") else "unresolved"
    if field == "snapshot_join_status":
        return row.get("snapshot_join_status", "source_only")
    return None


def _project(rows: Iterable[dict[str, Any]], fields: list[str]) -> list[dict[str, Any]]:
    return [{field: _value(row, field) for field in fields} for row in rows]


def _number(value: Any) -> float:
    try:
        return float(value) if value not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def _group_sum(rows: list[dict[str, Any]], keys: list[str], sums: list[str]) -> list[dict[str, Any]]:
    groups: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        key = tuple(_value(row, field) for field in keys)
        target = groups.setdefault(key, {field: value for field, value in zip(keys, key)})
        for field in sums:
            target[field] = _number(target.get(field)) + _number(_value(row, field))
    return list(groups.values())


def _advertising_daily(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    keys = ["advertiser_id", "advertiser_name", "metric_date", "ad_currency"]
    sums = ["ad_cost", "ad_gross_revenue", "sns_opt_gross_revenue", "impressions", "clicks", "sku_orders"]
    grouped = _group_sum(rows, keys, sums)
    by_key: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        by_key[tuple(_value(row, key) for key in keys)].append(row)
    for result in grouped:
        members = by_key[tuple(result[key] for key in keys)]
        result.update({
            "campaign_count_with_data": len({str(_value(row, "campaign_id")) for row in members}),
            "campaign_count_with_spend": len({str(_value(row, "campaign_id")) for row in members if _number(_value(row, "ad_cost")) > 0}),
            "campaign_count_with_orders": len({str(_value(row, "campaign_id")) for row in members if _number(_value(row, "sku_orders")) > 0}),
            "campaign_count_with_spend_zero_revenue": len({str(_value(row, "campaign_id")) for row in members if _number(_value(row, "ad_cost")) > 0 and _number(_value(row, "ad_gross_revenue")) == 0}),
        })
    return grouped


def _order_daily(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    orders: dict[str, dict[str, Any]] = {}
    for row in rows:
        order_id = str(row.get("order_id") or "")
        if not order_id:
            continue
        head = orders.setdefault(order_id, dict(row))
        for field in ("quantity", "returned_sku_quantity", "sku_subtotal_before_discount", "sku_seller_discount", "sku_platform_discount", "shipping_fee_after_discount", "shipping_fee_seller_discount", "shipping_fee_platform_discount", "taxes", "order_refund_amount"):
            head[f"_sum_{field}"] = _number(head.get(f"_sum_{field}")) + _number(row.get(field))
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for order in orders.values():
        day = str(order.get("created_at") or "")[:10]
        grouped[(day, str(order.get("currency") or ""))].append(order)
    results = []
    for (day, currency), members in grouped.items():
        statuses = Counter(str(item.get("order_status") or "").lower() for item in members)
        paid = sum(bool(item.get("paid_at")) for item in members)
        total_amount = sum(_number(item.get("order_amount")) for item in members)
        result = {
            "metric_date": day, "commerce_currency": currency, "distinct_order_count": len(members), "paid_order_count": paid,
            "current_completed_order_count": sum("complete" in status for status in statuses.elements()),
            "current_shipped_order_count": sum("ship" in status for status in statuses.elements()),
            "current_cancelled_order_count": sum("cancel" in status for status in statuses.elements()),
            "total_quantity": sum(_number(item.get("_sum_quantity")) for item in members),
            "returned_sku_quantity": sum(_number(item.get("_sum_returned_sku_quantity")) for item in members),
            "order_amount_distinct": total_amount, "average_order_amount": safe_div(total_amount, len(members)),
            "sku_subtotal_before_discount": sum(_number(item.get("_sum_sku_subtotal_before_discount")) for item in members),
            "seller_product_discount": sum(_number(item.get("_sum_sku_seller_discount")) for item in members),
            "platform_product_discount": sum(_number(item.get("_sum_sku_platform_discount")) for item in members),
            "shipping_fee_after_discount": sum(_number(item.get("_sum_shipping_fee_after_discount")) for item in members),
            "seller_shipping_discount": sum(_number(item.get("_sum_shipping_fee_seller_discount")) for item in members),
            "platform_shipping_discount": sum(_number(item.get("_sum_shipping_fee_platform_discount")) for item in members),
            "tax_amount": sum(_number(item.get("_sum_taxes")) for item in members),
            "order_refund_amount": sum(_number(item.get("_sum_order_refund_amount")) for item in members),
            "status_as_of_utc": max((str(item.get("captured_at_utc") or "") for item in members), default=""),
        }
        known = result["current_completed_order_count"] + result["current_shipped_order_count"] + result["current_cancelled_order_count"]
        result["other_status_order_count"] = len(members) - known
        result["unknown_order_status_count"] = sum(not str(item.get("order_status") or "").strip() for item in members)
        results.append(result)
    return results


def _returns_daily(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[(str(row.get("after_sale_applied_at") or "")[:10], str(row.get("currency") or ""))].append(row)
    results = []
    for (day, currency), members in groups.items():
        types = Counter(str(row.get("after_sale_type") or "unknown") for row in members)
        statuses = Counter(str(row.get("processing_status") or "unknown") for row in members)
        logistics = Counter(str(row.get("return_logistics_status") or "unknown") for row in members)
        results.append({
            "metric_date": day, "as_of_utc": max((str(row.get("captured_at") or "") for row in members), default=""),
            "commerce_currency": currency, "after_sale_case_count": len({str(row.get("after_sale_id")) for row in members}),
            "affected_order_count": len({str(row.get("order_id")) for row in members}),
            "refund_amount_current_known": sum(_number(row.get("refund_amount")) for row in members),
            "refund_only_case_count": sum("refund" in key.lower() and "return" not in key.lower() for key in types.elements()),
            "return_refund_case_count": sum("return" in key.lower() for key in types.elements()),
            "open_case_count": sum(any(word in key.lower() for word in ("open", "processing", "pending")) for key in statuses.elements()),
            "closed_case_count": sum(any(word in key.lower() for word in ("closed", "complete", "success")) for key in statuses.elements()),
            "status_mapping_unknown_count": statuses.get("unknown", 0),
            "after_sale_type_distribution_json": json.dumps(types, ensure_ascii=False, sort_keys=True),
            "processing_status_distribution_json": json.dumps(statuses, ensure_ascii=False, sort_keys=True),
            "return_logistics_status_distribution_json": json.dumps(logistics, ensure_ascii=False, sort_keys=True),
        })
    return results


def _join_campaign_status(datasets: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    intervals = {(str(row.get("advertiser_id")), str(row.get("campaign_id"))): row for row in datasets.get("gmvmax_campaign_performance_interval", [])}
    results = []
    for config in datasets.get("gmvmax_campaign_current_configuration_snapshot", []):
        interval = intervals.get((str(config.get("advertiser_id")), str(config.get("campaign_id"))))
        results.append({**(interval or {}), **config, "snapshot_join_status": "matched" if interval else "configuration_only"})
    return results


def _join_creative_authorization(datasets: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    authorization: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in datasets.get("creative_product_interval", []):
        authorization[(str(row.get("product_id") or ""), str(row.get("video_id") or ""))].append(row)
    results = []
    for row in datasets.get("gmvmax_campaign_product_creative_daily", []):
        key = (str(row.get("product_id") or ""), str(row.get("post_id") or ""))
        matches = authorization.get(key, [])
        auth = matches[0] if len(matches) == 1 else {}
        results.append({
            **row, **{field: auth.get(field) for field in ("authorization_type", "authorization_status", "pre_review_status")},
            "authorization_checked_at_utc": auth.get("captured_at_utc"),
            "authorization_match_status": "matched" if len(matches) == 1 else "ambiguous" if len(matches) > 1 else "unmatched",
        })
    return results


def _join_inventory(datasets: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    sku_status = {str(row.get("sku_id")): row for row in datasets.get("shop_sku_commercial_status_snapshot", [])}
    product_status = {str(row.get("product_id")): row for row in datasets.get("shop_product_commercial_status_snapshot", [])}
    results = []
    for inventory in datasets.get("inventory_snapshot_fact", []):
        sku = sku_status.get(str(inventory.get("sku_id")))
        product = product_status.get(str(inventory.get("product_id")))
        join_parts = [name for name, value in (("inventory", inventory), ("sku", sku), ("product", product)) if value]
        results.append({**(product or {}), **(sku or {}), **inventory, "snapshot_join_status": "+".join(join_parts)})
    return results


def _risk_metrics(datasets: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    rows = []
    for source, domain in (("shop_health_metric_snapshot_fact", "店铺健康"), ("fulfillment_performance_snapshot_fact", "履约")):
        for item in datasets.get(source, []):
            current = item.get("current_value_numeric", item.get("metric_value_official"))
            operator = item.get("target_operator")
            target = item.get("target_value_numeric")
            breach = None
            if current not in (None, "") and target not in (None, "") and operator in {"<", "<=", ">", ">="}:
                left, right = float(current), float(target)
                breach = {"<": left >= right, "<=": left > right, ">": left <= right, ">=": left < right}[operator]
            rows.append({
                **item, "risk_domain": domain, "metric_code": hashlib.sha256(f"{source}|{item.get('category')}|{item.get('metric_name')}".encode()).hexdigest()[:16],
                "metric_name_official": item.get("metric_name"), "as_of_utc": item.get("captured_at"),
                "current_value_raw": item.get("current_value_raw", item.get("metric_value_official")),
                "current_value_numeric": current, "current_value_unit": item.get("current_value_unit", item.get("unit")),
                "usable_value_numeric": current, "breach_status": "breached" if breach else "within_target" if breach is False else "unknown",
                "metric_mapping_status": "mapped" if current not in (None, "") else "raw_only",
            })
    return rows


def _channel_daily(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    channels = (("merchant_live", "商家直播", "overall_merchant_live_gmv"), ("merchant_video", "商家视频", "overall_merchant_video_gmv"), ("creator_affiliate", "达人联盟", "overall_creator_attributed_gmv"), ("merchant_product_card", "商家商品卡", "overall_merchant_product_card_gmv"))
    groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[(str(row.get("metric_date")), str(row.get("currency")))].append(row)
    results = []
    for (day, currency), members in groups.items():
        total = sum(_number(row.get("overall_gmv")) for row in members)
        values = {code: sum(_number(row.get(field)) for row in members) for code, _label, field in channels}
        delta = total - sum(values.values())
        for code, label, _field in channels:
            results.append({"metric_date": day, "channel_code": code, "channel_name": label, "commerce_currency": currency, "channel_gmv": values[code], "total_shop_gmv": total, "channel_gmv_share": safe_div(values[code], total), "reconciliation_delta": delta, "reconciliation_status": "matched" if abs(delta) < 0.01 else "mismatch"})
    return results


def _product_contribution(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[(str(row.get("metric_date")), str(row.get("currency")))].append(row)
    results = []
    for (_day, _currency), members in groups.items():
        total_gmv = sum(_number(row.get("overall_gmv")) for row in members)
        total_orders = sum(_number(row.get("overall_sku_orders")) for row in members)
        ranked = sorted(members, key=lambda row: (-_number(row.get("overall_gmv")), str(row.get("product_id"))))
        for rank, row in enumerate(ranked, 1):
            results.append({**row, "product_gmv": row.get("overall_gmv"), "gmv_share": safe_div(row.get("overall_gmv"), total_gmv), "gmv_rank": rank, "product_attributed_orders": row.get("overall_orders"), "sku_orders": row.get("overall_sku_orders"), "sku_order_share": safe_div(row.get("overall_sku_orders"), total_orders), "units_sold": row.get("overall_units_sold"), "impressions": row.get("overall_impressions"), "clicks": row.get("overall_clicks"), "ctr_ratio": safe_div(row.get("overall_clicks"), row.get("overall_impressions")), "refund_amount": row.get("overall_refund_amount"), "refunded_units": row.get("overall_refunded_units")})
    return results


def _advertised_product_contribution(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    keys = ["advertiser_id", "advertiser_name", "metric_date", "ad_currency", "product_id", "product_name_as_of"]
    grouped = _group_sum(rows, keys, ["ad_cost", "ad_gross_revenue", "sku_orders"])
    partitions: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in grouped:
        partitions[(str(row["advertiser_id"]), str(row["metric_date"]), str(row["ad_currency"]))].append(row)
    for members in partitions.values():
        total_cost = sum(_number(row.get("ad_cost")) for row in members); total_revenue = sum(_number(row.get("ad_gross_revenue")) for row in members)
        cost_rank = {id(row): rank for rank, row in enumerate(sorted(members, key=lambda item: -_number(item.get("ad_cost"))), 1)}
        revenue_rank = {id(row): rank for rank, row in enumerate(sorted(members, key=lambda item: -_number(item.get("ad_gross_revenue"))), 1)}
        for row in members:
            row.update({"cost_share": safe_div(row.get("ad_cost"), total_cost), "revenue_share": safe_div(row.get("ad_gross_revenue"), total_revenue), "cost_rank": cost_rank[id(row)], "revenue_rank": revenue_rank[id(row)]})
    return grouped


def _inventory_product_risk(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows: groups[str(row.get("product_id"))].append(row)
    results = []
    for product_id, members in groups.items():
        active = [row for row in members if str(_value(row, "listing_status") or "").lower() not in {"deleted", "delisted", "inactive"}]
        out = [row for row in active if _number(row.get("available_stock")) == 0]
        low = [row for row in active if str(row.get("inventory_status_calculated") or row.get("inventory_status_raw") or "").lower() == "low_stock"]
        finite = [_number(row.get("days_of_supply_calculated")) for row in active if row.get("days_of_supply_calculated") not in (None, "")]
        prices = [_number(row.get("effective_price_calculated")) for row in active if row.get("effective_price_calculated") not in (None, "")]
        risk = "out_of_stock" if out else "low_stock" if low else "normal" if active else "unknown"
        reasons = (["存在可售库存为0的活跃SKU"] if out else []) + (["存在低库存SKU"] if low else [])
        first = members[0]
        results.append({**first, "product_id": product_id, "product_name": first.get("product_name"), "inventory_as_of_utc": first.get("captured_at"), "commercial_status_as_of_utc": first.get("captured_at_utc"), "commerce_currency": first.get("currency"), "listing_status": _value(first, "listing_status"), "sku_count": len(members), "active_sku_count": len(active), "out_of_stock_sku_count": len(out), "low_stock_sku_count": len(low), "total_stock": sum(_number(row.get("total_stock")) for row in members), "available_stock": sum(_number(row.get("available_stock")) for row in members), "locked_stock": sum(_number(row.get("locked_stock")) for row in members), "sales_30d": sum(_number(row.get("sales_30d")) for row in members), "forecast_30d": sum(_number(row.get("forecast_30d")) for row in members), "recommended_replenishment": sum(_number(row.get("recommended_replenishment")) for row in members), "min_finite_days_of_supply": min(finite) if finite else None, "min_effective_price": min(prices) if prices else None, "max_effective_price": max(prices) if prices else None, "promotion_active_sku_count": sum(bool(row.get("sku_level_promotion_ids_json") not in (None, "", "[]")) for row in active), "inventory_risk_status": risk, "risk_reasons_json": json.dumps(reasons, ensure_ascii=False)})
    return results


def _anomalies(datasets: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    results = []
    for row in _inventory_product_risk(_join_inventory(datasets)):
        if row.get("inventory_risk_status") not in {"out_of_stock", "low_stock"}:
            continue
        evidence = f"商品可售库存 {row.get('available_stock')}，风险状态 {row.get('inventory_risk_status')}"
        seed = f"inventory|{row.get('product_id')}|{row.get('inventory_as_of_utc')}|{row.get('inventory_risk_status')}"
        results.append({
            "anomaly_id": hashlib.sha256(seed.encode()).hexdigest()[:24], "detected_at_utc": row.get("inventory_as_of_utc"),
            "effective_date": str(row.get("inventory_as_of_utc") or "")[:10], "evidence_as_of_utc": row.get("inventory_as_of_utc"),
            "entity_type": "product", "entity_id": row.get("product_id"), "entity_name": row.get("product_name"),
            "anomaly_domain": "inventory", "anomaly_code": row.get("inventory_risk_status"),
            "severity": "high" if row.get("inventory_risk_status") == "out_of_stock" else "medium",
            "observed_value": row.get("available_stock"), "observed_unit": "件", "baseline_type": "rule",
            "threshold_operator": "=", "threshold_value": 0 if row.get("inventory_risk_status") == "out_of_stock" else None,
            "evidence_summary": evidence, "evidence_table": "SKU库存当前状态", "evidence_primary_key": row.get("product_id"),
            "candidate_reasons_json": row.get("risk_reasons_json"), "evidence_completeness_status": "complete",
            "diagnosis_status": "fact_only", "rule_version": DASHBOARD_CONTRACT_VERSION, "resolution_status": "open",
        })
    return results


def _owner_anomalies(datasets: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    rows = []
    for item in _anomalies(datasets):
        rows.append({
            "risk_id": item["anomaly_id"], "detected_at_utc": item["detected_at_utc"], "evidence_as_of_utc": item["evidence_as_of_utc"],
            "risk_domain": item["anomaly_domain"], "severity": item["severity"], "risk_name": item["anomaly_code"],
            "entity_type": item["entity_type"], "entity_id": item["entity_id"], "entity_name": item["entity_name"],
            "observed_value": item["observed_value"], "observed_unit": item["observed_unit"], "observed_currency": item.get("observed_currency"),
            "threshold_operator": item.get("threshold_operator"), "threshold_value": item.get("threshold_value"),
            "evidence_summary": item["evidence_summary"], "decision_required": "确认补货或调整商品投放",
            "evidence_table": item["evidence_table"], "evidence_primary_key": item["evidence_primary_key"],
            "evidence_completeness_status": item["evidence_completeness_status"], "rule_version": item["rule_version"], "resolution_status": item["resolution_status"],
        })
    return rows


def _rows_for(table: dict[str, Any], datasets: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    name = table["table_name"]
    source = table["sources"][0]
    rows = datasets.get(source, [])
    if name == "计划当前状态":
        rows = _join_campaign_status(datasets)
    elif name == "计划素材日表现":
        rows = _join_creative_authorization(datasets)
    elif name == "SKU库存当前状态":
        rows = _join_inventory(datasets)
    elif name == "店铺承接风险当前状态":
        rows = _risk_metrics(datasets)
    elif name == "广告经营日报":
        rows = _advertising_daily(rows)
    elif name == "订单日报":
        rows = _order_daily(rows)
    elif name == "售后申请日报":
        rows = _returns_daily(rows)
    elif name == "店铺经营日报":
        products = _group_sum(datasets.get("shop_product_performance_daily", []), ["metric_date", "commerce_currency"], ["product_gmv", "sku_orders", "units_sold", "estimated_customers_product", "impressions", "unique_impressions"])
        orders = {(row["metric_date"], row["commerce_currency"]): row for row in _order_daily(datasets.get("order_line_operations", []))}
        rows = []
        for product in products:
            order = orders.get((product["metric_date"], product["commerce_currency"]), {})
            rows.append({**product, "shop_gmv": product.get("product_gmv"), "orders_created": order.get("distinct_order_count"), "order_amount_created": order.get("order_amount_distinct"), "average_order_amount": order.get("average_order_amount"), "customers_official": product.get("estimated_customers_product"), "product_impressions_official": product.get("impressions"), "unique_product_impressions_official": product.get("unique_impressions"), "product_data_status":"complete", "order_data_status":"complete" if order else "unavailable", "overview_data_status":"unavailable"})
    elif name == "渠道贡献日报":
        rows = _channel_daily(rows)
    elif name == "商品经营贡献日报":
        rows = _product_contribution(rows)
    elif name == "商品广告贡献日报":
        rows = _advertised_product_contribution(rows)
    elif name == "库存风险当前状态":
        rows = _inventory_product_risk(_join_inventory(datasets))
    elif name == "异常事实":
        rows = _anomalies(datasets)
    elif name == "经营异常清单":
        rows = _owner_anomalies(datasets)
    return _project(rows, table["fields"])


def _coerce(value: Any) -> Any:
    if value in (None, ""):
        return None
    text = str(value)
    if re.fullmatch(r"-?\d+", text):
        try: return int(text)
        except ValueError: pass
    if re.fullmatch(r"-?(?:\d+\.\d*|\d*\.\d+)", text):
        try: return float(text)
        except ValueError: pass
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        try: return datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError: pass
    return value


def _content_hash(fields: list[str], rows: list[dict[str, Any]]) -> str:
    normalized = [[_coerce(row.get(field)) for field in fields] for row in rows]
    return hashlib.sha256(json.dumps(normalized, ensure_ascii=False, sort_keys=False, separators=(",", ":"), default=str).encode("utf-8")).hexdigest()


def build_dashboard_workbooks(staging: Path, datasets: dict[str, list[dict[str, Any]]], contract_path: Path, official_no_data_sources: set[str] | None = None) -> list[dict[str, Any]]:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise DashboardBuildError("dashboard_xlsx_runtime_missing: openpyxl==3.1.5") from exc
    contract = json.loads(contract_path.read_text(encoding="utf-8-sig"))
    official_no_data_sources = official_no_data_sources or set()
    output_root = staging / "dashboard_tables"
    entries = []
    for table in contract["tables"]:
        required = [source for source in table["sources"] if source != "derived"]
        missing = [source for source in required if source not in datasets]
        if missing and not all(source in official_no_data_sources for source in missing):
            entries.append({"audience": table["audience"], "table_name": table["table_name"], "status": "source_missing", "missing_sources": missing, "explicit_only": bool(table.get("explicit_only"))})
            continue
        fields = table["fields"]
        effective_datasets = dict(datasets)
        for source in missing:
            effective_datasets[source] = []
        rows = _rows_for(table, effective_datasets)
        target = output_root / table["audience"] / f"{table['table_name']}.xlsx"
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = table["table_name"][:31]
        sheet.freeze_panes = "A2"
        sheet.append([visible_header(field) for field in fields])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            sheet.append([_coerce(row.get(field)) for field in fields])
        sheet.auto_filter.ref = sheet.dimensions
        ratio_columns = {index + 1 for index, field in enumerate(fields) if field.endswith("_ratio") or field.endswith("_share") or field in {"roi", "sns_opt_roi", "target_roi", "roas_bid"}}
        money_columns = {index + 1 for index, field in enumerate(fields) if any(token in field for token in ("cost", "revenue", "gmv", "amount", "price", "budget", "discount", "fee", "tax"))}
        for column in range(1, len(fields) + 1):
            values = [str(sheet.cell(row=row, column=column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(42, max(10, max(map(len, values), default=10) + 2))
            number_format = "0.00%" if column in ratio_columns else "#,##0.00" if column in money_columns else None
            if number_format:
                for cell in sheet[get_column_letter(column)][1:]:
                    cell.number_format = number_format
        _save_deterministic_xlsx(workbook, target)
        checked = load_workbook(target, read_only=False, data_only=True)
        checked_sheet = checked.active
        if [cell.value for cell in checked_sheet[1]] != [visible_header(field) for field in fields] or checked_sheet.max_row != len(rows) + 1:
            raise DashboardBuildError(f"dashboard_xlsx_verification_failed: {table['table_name']}")
        checked.close()
        entries.append({
            "audience": table["audience"], "table_name": table["table_name"], "status": "official_zero" if not rows else "generated",
            "relative_path": target.relative_to(staging).as_posix(), "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
            "business_content_sha256": _content_hash(fields, rows), "row_count": len(rows), "fields": fields,
            "source_datasets": table["sources"], "grain": table["grain"], "explicit_only": bool(table.get("explicit_only")),
        })
    return entries
