"""Incrementally publish authoritative artifacts into a shop-facing directory."""
from __future__ import annotations

import csv
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import tempfile
from typing import Any


class BusinessDirectoryError(RuntimeError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _fail(code: str, message: str) -> None:
    raise BusinessDirectoryError(code, message)


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise BusinessDirectoryError("business_directory_manifest_invalid", str(path)) from exc
    if not isinstance(value, dict):
        _fail("business_directory_manifest_invalid", str(path))
    return value


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _xlsx_content_hash(path: Path) -> str:
    runtime = Path(__file__).resolve().parents[1] / ".runtime"
    if runtime.is_dir() and str(runtime) not in os.sys.path:
        os.sys.path.insert(0, str(runtime))
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise BusinessDirectoryError("business_directory_xlsx_runtime_missing", str(path)) from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
    workbook.close()
    def canonical(value: Any) -> Any:
        if hasattr(value, "isoformat"):
            text = value.isoformat()
            return text[:10] if text.endswith("T00:00:00") else text
        return value
    payload = json.dumps([[canonical(value) for value in row] for row in values], ensure_ascii=False, separators=(",", ":"), default=str).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _verify(root: Path, relative: object, expected: object) -> Path:
    text = str(relative or "")
    if not text or "\\" in text or any(part in {"", ".", ".."} for part in text.split("/")):
        _fail("business_directory_source_path_invalid", text)
    source = (root / Path(*text.split("/"))).resolve()
    try:
        source.relative_to(root.resolve())
    except ValueError:
        _fail("business_directory_source_path_invalid", text)
    expected_text = str(expected or "")
    if not source.is_file() or not re.fullmatch(r"[a-f0-9]{64}", expected_text):
        _fail("business_directory_source_invalid", text)
    if _sha256(source) != expected_text:
        _fail("business_directory_source_hash_mismatch", text)
    return source


def _segment(value: object, fallback: str) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value or "").strip()).rstrip(" .")
    return (text or fallback)[:120]


def _captured_date(value: object) -> str | None:
    match = re.match(r"\d{4}-\d{2}-\d{2}", str(value or ""))
    return match.group(0) if match else None


def _date_label(request: dict[str, Any], captured: object, time_naming: object) -> str:
    if str(time_naming or "") == "snapshot":
        return _captured_date(captured) or "当前快照"
    start = str(request.get("start_date") or "")
    end = str(request.get("end_date") or "")
    if start and end:
        return start if start == end else f"{start}_至_{end}"
    return _captured_date(captured) or "当前快照"


def _load_mapping(path: Path) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    value = _read_json(path)
    return (
        {str(item["report_id"]): item for item in value.get("reports", [])},
        {str(item["dataset_name"]): item for item in value.get("datasets", [])},
    )


def _raw_items(export_root: Path, shop_id: str, mappings: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], set[str]]:
    items: list[dict[str, Any]] = []
    aliases: set[str] = set()
    for manifest_path in sorted((export_root / "batches").glob("*/manifest.json")):
        manifest = _read_json(manifest_path)
        shop = manifest.get("shop") or {}
        if str(shop.get("shop_id") or shop.get("seller_id") or "") != shop_id:
            continue
        alias = str(shop.get("shop_alias") or "").strip()
        if alias:
            aliases.add(alias)
        request = manifest.get("request") or {}
        for artifact in manifest.get("artifacts") or []:
            name = str(artifact.get("report_id") or "")
            mapping = mappings.get(name)
            if mapping is None:
                continue
            artifact_relative = str(artifact.get("relative_path") or "")
            relative = (manifest_path.parent.relative_to(export_root) / Path(*artifact_relative.split("/"))).as_posix()
            source = _verify(export_root, relative, artifact.get("sha256"))
            captured = artifact.get("captured_at_utc") or manifest.get("published_at_utc") or manifest.get("created_at")
            scope_identity = artifact.get("scope_identity") or {}
            source_metadata = artifact.get("source_metadata") or {}
            identity = json.dumps(scope_identity, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            items.append({
                "layer": "RAW", "name": name, "mapping": mapping, "source": source,
                "sha256": artifact.get("sha256"), "date_label": _date_label(request, captured, mapping.get("time_naming")),
                "identity": identity, "batch_id": manifest.get("export_batch_id"),
                "published_at": manifest.get("published_at_utc") or manifest.get("created_at_utc") or manifest.get("created_at") or "",
                "source_relative_path": relative, "source_kind": mapping.get("source_kind"),
                "scope_identity": scope_identity, "source_metadata": source_metadata,
            })
    return items, aliases


def _openpyxl() -> tuple[Any, Any]:
    runtime = Path(__file__).resolve().parents[1] / ".runtime"
    if runtime.is_dir() and str(runtime) not in os.sys.path:
        os.sys.path.insert(0, str(runtime))
    try:
        from openpyxl import Workbook, load_workbook
    except ModuleNotFoundError as exc:
        raise BusinessDirectoryError("business_directory_xlsx_runtime_missing", "openpyxl") from exc
    return Workbook, load_workbook


def _merge_creative_batch(items: list[dict[str, Any]], staging: Path) -> dict[str, Any]:
    Workbook, load_workbook = _openpyxl()
    headers: list[Any] | None = None
    rows: list[list[Any]] = []
    source_paths: list[str] = []
    for item in sorted(items, key=lambda value: (
        str(value["source_metadata"].get("campaign_id") or ""),
        str(value["source_metadata"].get("product_id") or ""),
        str(value["source_relative_path"]),
    )):
        workbook = load_workbook(item["source"], read_only=True, data_only=True)
        sheet = workbook.active
        sheet.reset_dimensions()
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        if not values:
            _fail("business_directory_merge_source_empty", str(item["source"]))
        current_headers = values[0]
        if headers is None:
            headers = current_headers
        elif current_headers != headers:
            _fail("business_directory_merge_header_mismatch", str(item["source"]))
        metadata = item["source_metadata"]
        prefix = [metadata.get("campaign_id"), metadata.get("campaign_name_observed"), metadata.get("product_id"), metadata.get("product_name_observed")]
        rows.extend([prefix + row for row in values[1:] if any(value not in (None, "") for value in row)])
        source_paths.append(item["source_relative_path"])
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["计划ID", "计划名称", "商品ID", "商品名称", *(headers or [])])
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    target = staging / f"creative-{items[0]['batch_id']}.xlsx"
    workbook.save(target)
    workbook.close()
    merged = dict(items[0])
    merged.update({
        "source": target,
        "sha256": _sha256(target),
        "content_sha256": _xlsx_content_hash(target),
        "identity": json.dumps({"advertiser_id": items[0]["scope_identity"].get("advertiser_id")}, sort_keys=True, separators=(",", ":")),
        "source_relative_path": json.dumps(source_paths, ensure_ascii=False, separators=(",", ":")),
        "source_kind": "merged_official_partitions_xlsx",
        "row_count": len(rows),
    })
    return merged


def _render_product_daily_api(item: dict[str, Any], staging: Path) -> dict[str, Any]:
    Workbook, _ = _openpyxl()
    envelope = _read_json(item["source"])
    if envelope.get("schema_version") != "seller_shop_product_daily_api_envelope_v1":
        _fail("business_directory_product_api_invalid", str(item["source"]))
    rows: list[dict[str, Any]] = []
    metric_keys: set[str] = set()
    for collection in envelope.get("collections") or []:
        metric_date = collection.get("metric_date")
        for channel in collection.get("channels") or []:
            channel_name = channel.get("channel_key")
            for product in channel.get("items") or []:
                meta = product.get("meta") or {}
                stats = (product.get("stats_v3") or {}).get(channel_name) or {}
                values: dict[str, Any] = {}
                for key, value in stats.items():
                    if isinstance(value, dict) and value.get("amount") is not None:
                        values[key] = value.get("amount")
                    elif not isinstance(value, (dict, list)):
                        values[key] = value
                metric_keys.update(values)
                rows.append({
                    "日期": metric_date, "商品ID": str(meta.get("product_id") or ""),
                    "商品名称": meta.get("product_name"), "GMV层级代码": meta.get("gmv_range"),
                    "商品状态代码": meta.get("product_status"), "渠道": channel_name, **values,
                })
    fixed = ["日期", "商品ID", "商品名称", "GMV层级代码", "商品状态代码", "渠道"]
    headers = fixed + sorted(metric_keys)
    workbook = Workbook(write_only=False)
    sheet = workbook.active
    sheet.title = "商品每日表现"
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    target = staging / f"product-daily-{item['batch_id']}.xlsx"
    workbook.save(target)
    workbook.close()
    rendered = dict(item)
    rendered.update({
        "source": target, "sha256": _sha256(target), "content_sha256": _xlsx_content_hash(target),
        "source_kind": "official_api_snapshot_rendered_xlsx", "row_count": len(rows),
    })
    return rendered


def _merge_partitioned_raw(items: list[dict[str, Any]], staging: Path) -> list[dict[str, Any]]:
    passthrough: list[dict[str, Any]] = []
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for item in items:
        if item["name"] == "shop_product_daily_performance":
            if item["source"].suffix.lower() != ".json":
                _fail("business_directory_product_daily_source_unsupported", str(item["source"]))
            passthrough.append(_render_product_daily_api(item, staging))
            continue
        if item["name"] != "gmvmax_campaign_product_creative_performance" or item["source"].suffix.lower() != ".xlsx":
            passthrough.append(item)
            continue
        advertiser_id = str(item["scope_identity"].get("advertiser_id") or "")
        grouped.setdefault((str(item["batch_id"]), item["date_label"], advertiser_id), []).append(item)
    for group in grouped.values():
        passthrough.append(_merge_creative_batch(group, staging))
    return passthrough


def _clean_items(cleaning_root: Path, shop_id: str, mappings: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], set[str]]:
    items: list[dict[str, Any]] = []
    aliases: set[str] = set()
    for manifest_path in sorted((cleaning_root / "batches").glob("*/cleaning_manifest.json")):
        manifest = _read_json(manifest_path)
        shop = manifest.get("shop") or {}
        if str(shop.get("shop_id") or "") != shop_id or manifest.get("qa_status") != "passed":
            continue
        alias = str(shop.get("shop_alias") or "").strip()
        if alias:
            aliases.add(alias)
        request = manifest.get("request") or {}
        advertiser_id = str((manifest.get("selected_advertiser") or {}).get("advertiser_id") or "")
        for output in manifest.get("outputs") or []:
            name = str(output.get("dataset_name") or "")
            mapping = mappings.get(name)
            if mapping is None:
                continue
            source = _verify(cleaning_root, output.get("relative_path"), output.get("sha256"))
            captured = output.get("captured_at_utc") or manifest.get("published_at_utc")
            items.append({
                "layer": "CLEAN", "name": name, "mapping": mapping, "source": source,
                "sha256": output.get("sha256"), "date_label": _date_label(request, captured, mapping.get("time_naming")),
                "identity": advertiser_id, "batch_id": manifest.get("cleaning_batch_id"),
                "published_at": manifest.get("published_at_utc") or manifest.get("created_at_utc") or "",
                "source_relative_path": output.get("relative_path"), "source_kind": "cleaned_csv",
                "row_count": output.get("row_count"),
            })
        for output in manifest.get("dashboard_outputs") or []:
            if output.get("status") not in {"generated", "official_zero"}:
                continue
            source = _verify(cleaning_root, output.get("relative_path"), output.get("sha256"))
            mapping = {"chinese_group": output.get("audience"), "chinese_table_name": output.get("table_name"), "time_naming": "snapshot" if "当前状态" in str(output.get("table_name")) else "range"}
            captured = manifest.get("published_at_utc")
            items.append({
                "layer": "CLEAN", "name": f"dashboard:{output.get('audience')}:{output.get('table_name')}",
                "mapping": mapping, "source": source, "sha256": output.get("sha256"),
                "content_sha256": output.get("business_content_sha256"),
                "date_label": _date_label(request, captured, mapping.get("time_naming")),
                "identity": advertiser_id, "batch_id": manifest.get("cleaning_batch_id"),
                "published_at": manifest.get("published_at_utc") or manifest.get("created_at_utc") or "",
                "source_relative_path": output.get("relative_path"), "source_kind": "dashboard_xlsx",
                "row_count": output.get("row_count"),
            })
    return items, aliases


def _atomic_copy(source: Path, target: Path, expected_hash: str) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(prefix=".publish-", suffix=target.suffix, dir=target.parent)
    os.close(handle)
    temporary = Path(temporary_name)
    try:
        shutil.copyfile(source, temporary)
        if _sha256(temporary) != expected_hash:
            _fail("business_directory_copy_hash_mismatch", str(target))
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def _target_for(item: dict[str, Any], shop_root: Path) -> tuple[Path, int, bool, str]:
    mapping = item["mapping"]
    layer = "10_导出原始" if item["layer"] == "RAW" else "20_清洗标准"
    table = _segment(mapping.get("chinese_table_name"), item["name"])
    directory = shop_root / layer / _segment(mapping.get("chinese_group"), "未分类") / table
    suffix = item["source"].suffix.lower()
    version = 1
    while True:
        marker = "" if version == 1 else f"_v{version}"
        target = directory / f"{item['date_label']}{marker}{suffix}"
        if not target.exists():
            return target, version, False, str(item["sha256"])
        if target.is_file():
            target_sha256 = _sha256(target)
            if (
                target_sha256 == item["sha256"]
                or (item.get("content_sha256") and target.suffix.lower() == ".xlsx" and _xlsx_content_hash(target) == item["content_sha256"])
            ):
                return target, version, True, target_sha256
        version += 1


def _publish_item(item: dict[str, Any], shop_root: Path) -> tuple[Path, int, bool, str]:
    target, version, was_reused, target_sha256 = _target_for(item, shop_root)
    if not was_reused:
        _atomic_copy(item["source"], target, item["sha256"])
    return target, version, was_reused, target_sha256


def _atomic_write(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(prefix=".metadata-", suffix=path.suffix, dir=path.parent)
    os.close(handle)
    temporary = Path(temporary_name)
    try:
        temporary.write_bytes(content)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _write_metadata(shop_root: Path, shop_id: str, shop_alias: str, rows: list[dict[str, Any]]) -> None:
    profile = shop_root / "00_店铺档案"
    records = shop_root / "99_生成记录_不用看"
    _atomic_write(profile / "店铺档案.csv", f"shop_id,shop_alias\n{shop_id},{shop_alias}\n".encode("utf-8-sig"))
    columns = ["layer", "report_or_dataset", "relative_path", "coverage", "sha256", "source_batch_id", "source_relative_path", "source_kind", "row_count", "version"]
    buffer = [",".join(columns)]
    for row in rows:
        buffer.append(",".join('"' + str(row.get(column, "")).replace('"', '""') + '"' for column in columns))
    _atomic_write(profile / "数据目录索引.csv", ("\ufeff" + "\n".join(buffer) + "\n").encode("utf-8"))
    index = {"schema_version": "business_directory_index_v3", "entries": rows}
    _atomic_write(profile / "数据目录索引.json", (json.dumps(index, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    manifest = {
        "schema_version": "business_directory_manifest_v3", "shop_id": shop_id, "shop_alias": shop_alias,
        "raw_file_count": sum(row["layer"] == "RAW" for row in rows),
        "clean_file_count": sum(row["layer"] == "CLEAN" for row in rows),
        "generated_files": [
            {"relative_path": row["relative_path"], "sha256": row["sha256"]} for row in rows
        ],
    }
    _atomic_write(records / "业务目录清单.json", (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))


def refresh_business_directory(*, export_root: Path, cleaning_root: Path, business_root: Path, shop_id: str, mapping_path: Path) -> dict[str, Any]:
    """Append missing authoritative files; never replace or delete a shop directory."""
    report_map, dataset_map = _load_mapping(mapping_path)
    raw, raw_aliases = _raw_items(export_root, shop_id, report_map)
    clean, clean_aliases = _clean_items(cleaning_root, shop_id, dataset_map)
    aliases = raw_aliases | clean_aliases
    if not raw or len(aliases) != 1:
        _fail("business_directory_shop_unresolved", f"Expected one shop alias for {shop_id}")
    shop_alias = next(iter(aliases))
    shop_root = business_root / _segment(shop_alias, "未知店铺")
    rows: list[dict[str, Any]] = []
    published = reused = 0
    with tempfile.TemporaryDirectory(prefix="tiktok-business-merge-") as staging_name:
        raw = _merge_partitioned_raw(raw, Path(staging_name))
        groups: dict[tuple[str, str, str, str, str], list[dict[str, Any]]] = {}
        for item in raw + clean:
            key = (item["layer"], item["name"], item["date_label"], item["identity"], item["source"].suffix.lower())
            groups.setdefault(key, []).append(item)
        for key in sorted(groups):
            seen: set[str] = set()
            for item in sorted(groups[key], key=lambda entry: (str(entry["published_at"]), str(entry["batch_id"]))):
                content_identity = str(item.get("content_sha256") or item["sha256"])
                if content_identity in seen:
                    continue
                seen.add(content_identity)
                target, version, was_reused, target_sha256 = _publish_item(item, shop_root)
                if was_reused:
                    reused += 1
                else:
                    published += 1
                rows.append({
                    "layer": item["layer"], "report_or_dataset": item["name"],
                    "relative_path": target.relative_to(shop_root).as_posix(), "coverage": item["date_label"],
                    "sha256": target_sha256, "source_batch_id": item["batch_id"],
                    "source_relative_path": item["source_relative_path"], "source_kind": item["source_kind"],
                    "row_count": item.get("row_count"), "version": version,
                    "business_content_sha256": item.get("content_sha256"),
                })
    _write_metadata(shop_root, shop_id, shop_alias, rows)
    return {
        "status": "published" if published else "reused", "shop_id": shop_id, "shop_alias": shop_alias,
        "output_dir": str(shop_root), "raw_file_count": sum(row["layer"] == "RAW" for row in rows),
        "clean_file_count": sum(row["layer"] == "CLEAN" for row in rows),
        "base_clean_file_count": sum(row["layer"] == "CLEAN" and row["source_kind"] == "cleaned_csv" for row in rows),
        "dashboard_file_count": sum(row["source_kind"] == "dashboard_xlsx" for row in rows),
        "published_file_count": published, "reused_file_count": reused,
    }
